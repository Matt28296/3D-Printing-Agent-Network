"""Write setup_page/cost_section.html (an itemized cost section) from the LibreOffice-recalculated workbook.
usage: python3 cost_section.py   (run recalc first: it leaves costs/rc/Robot build costs.xlsx)"""
import html, json, pathlib
from openpyxl import load_workbook

HERE = pathlib.Path(__file__).parent
wb = load_workbook(HERE / "rc" / "Robot build costs.xlsx", data_only=True)
prices = {}
for r in wb["Prices"].iter_rows(min_row=5, values_only=True):
    if r[0]:
        prices[r[0]] = {"retailer": r[3] or "", "listing": r[2] or ""}
summ = json.loads((HERE / "costs.json").read_text())
HST = 1.15
m = lambda v: f"${v:,.2f}"
e = lambda s: html.escape(str(s))


def robot(sheet, key, title):
    ws = wb[sheet]
    rows, used, buy, budget = [], 0.0, 0.0, 0.0
    for r in ws.iter_rows(min_row=5, values_only=True):
        group, item, pkey, qty, unit, val, packq, packs, packp, cost, bud = r[:11]
        if not item or group in (None, "") or isinstance(group, str) and group.startswith("Total"):
            continue
        if item in ("Total before tax", "HST", "Total with HST"):
            continue
        if not isinstance(val, (int, float)):
            continue
        retailer = prices.get(pkey, {}).get("retailer", "") if pkey else ""
        q = f"{qty:g}" if isinstance(qty, (int, float)) else e(qty)
        rows.append(f"<tr><td>{e(group)}</td><td>{e(item)}</td><td class='n'>{q}</td><td>{e(retailer)}</td>"
                    f"<td class='n'>{m(val)}</td><td class='n'>{m(cost or 0)}</td></tr>")
    s = summ[key]
    return (f"<details><summary><b>{e(title)}</b>: {m(s['scratch'])} buying every pack new · {m(s['used'])} for the parts it uses · "
            f"{m(s['budget'])} with budget picks (with HST)</summary>\n<div class='tbl'><table><thead><tr><th>Group</th><th>Item</th>"
            f"<th class='n'>Qty</th><th>Where (price checked)</th><th class='n'>Value used</th><th class='n'>Cost to buy</th></tr></thead><tbody>\n"
            + "\n".join(rows) + "\n</tbody></table></div></details>")


a = summ["all"]
out = f"""<section aria-labelledby="cost">
  <div class="sec-head"><h2 id="cost">What they cost</h2><span class="muted mono">CAD · prices checked {e(summ['asof'])} · 15 % HST</span></div>
  <p>Every bought part, from each robot's build guide, priced at Canadian retailers (a few US stores and labelled estimates where nothing Canadian could be read). <b>Value used</b> is the share of each pack the robot uses; <b>cost to buy</b> is whole packs for that robot alone. PETG and print electricity are included; shipping, tools and the Brio you already own are not. The full workbook, with links and formulas, is Robot build costs.xlsx in Google Drive and <code>docs/costs/</code> on GitHub.</p>
  <div class="tbl"><table><thead><tr><th>Robot</th><th class='n'>Buying every pack new</th><th class='n'>Parts it uses</th><th class='n'>Budget picks</th></tr></thead><tbody>
    <tr><td>Desktop robot arm</td><td class='n'>{m(summ['arm']['scratch'])}</td><td class='n'>{m(summ['arm']['used'])}</td><td class='n'>{m(summ['arm']['budget'])}</td></tr>
    <tr><td>Tracked rover</td><td class='n'>{m(summ['rover']['scratch'])}</td><td class='n'>{m(summ['rover']['used'])}</td><td class='n'>{m(summ['rover']['budget'])}</td></tr>
    <tr><td>Hexapod</td><td class='n'>{m(summ['hexapod']['scratch'])}</td><td class='n'>{m(summ['hexapod']['used'])}</td><td class='n'>{m(summ['hexapod']['budget'])}</td></tr>
    <tr><td><b>All three, bought together</b></td><td class='n'><b>{m(a['scratch'])}</b></td><td></td><td class='n'>{m(a['budget'])}</td></tr>
  </tbody></table></div>
  <p class="muted">Buying all three together saves about {m(a['saving'])}, because they share packs of screws, wire and small parts. Tap a robot below for its itemized list (values before tax).</p>
  {robot('Arm', 'arm', 'Desktop robot arm')}
  {robot('Rover', 'rover', 'Tracked rover')}
  {robot('Hexapod', 'hexapod', 'Hexapod')}
</section>
"""
(HERE.parent / "setup_page" / "cost_section.html").write_text(out)
print("cost section", len(out), "chars")
