"""Rebuild the cost workbook, recalculate it with LibreOffice, and write costs/costs.json with the totals.
usage: python3 recalc.py"""
import json, subprocess, sys, tempfile, pathlib
from openpyxl import load_workbook
HERE = pathlib.Path(__file__).parent
subprocess.run([sys.executable, str(HERE / "build_xlsx.py")], check=True)
with tempfile.TemporaryDirectory() as td:
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", td,
                    str(HERE / "Robot build costs.xlsx")], check=True, capture_output=True)
    ws = load_workbook(pathlib.Path(td) / "Robot build costs.xlsx", data_only=True)["Summary"]
    v = lambda c: float(ws[c].value)
    out = {"asof": "26 Oct 2026",
           "arm": {"scratch": v("G15"), "used": v("G23"), "budget": v("H15")},
           "rover": {"scratch": v("G16"), "used": v("G24"), "budget": v("H16")},
           "hexapod": {"scratch": v("G17"), "used": v("G25"), "budget": v("H17")},
           "all": {"scratch": v("G18"), "budget": v("H18"), "saving": v("G19")}}
(HERE / "costs.json").write_text(json.dumps(out, indent=1))
for k in ("arm", "rover", "hexapod"):
    print(f"{k:8s} scratch {out[k]['scratch']:7.2f}  used {out[k]['used']:7.2f}  budget {out[k]['budget']:7.2f}")
print(f"all three {out['all']['scratch']:.2f} (budget {out['all']['budget']:.2f}), saving {out['all']['saving']:.2f}")
