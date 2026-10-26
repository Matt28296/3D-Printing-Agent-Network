"""Build 'Robot build costs.xlsx' from the price lists and the per-robot bought-parts lists (bom.py).
Every total is a live formula; prices sit on the Prices sheet, quantities on each robot sheet."""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent  # prices_*.csv and bom.py live next to this script
sys.path.insert(0, str(HERE))
import bom  # noqa: E402

OUT = HERE / "Robot build costs.xlsx"
FX = 1.3993          # Bank of Canada, 23 Oct 2026 (price research A)
AS_OF = "26 Oct 2026"

# ---------------------------------------------------------------------------------------- price table
def load(name):
    p = HERE / name
    return list(csv.DictReader(open(p))) if p.exists() else []

main, budget = {}, {}
for r in load("prices_a.csv") + load("prices_b.csv") + load("prices_est.csv"):
    k = r["key"]
    is_budget = "(budget)" in r["retailer"] or r["notes"].startswith("BUDGET")
    if is_budget:
        budget[k] = r
        continue
    if r["confidence"] == "none":
        continue
    if k in main and main[k]["confidence"] in ("high", "medium") and r["confidence"] == "estimate":
        continue
    main[k] = r

# fixes and fills decided on review of the research
main["mp1584"] = dict(budget["mp1584"], notes="Only an AliExpress listing was found (no Canadian stock seen)")
main["cells18650"] = {"key": "cells18650", "listing": "Protected 18650 Li-ion cell, ≈3000-3500 mAh, button top",
                      "retailer": "Canadian battery stores (typical)", "url": "", "pack_qty": "1", "pack_price": "23.00",
                      "currency": "CAD", "confidence": "medium", "date_checked": "2026-10-26",
                      "notes": "Research found C$22-24 each in stock at Canadian stores; US store Illumn sells LG MJ1 protected at US$7.49 but Li-ion shipping to Canada is unconfirmed"}
budget["cells18650"] = [r for r in load("prices_a.csv") if r["key"] == "cells18650"][0]

NAMES = {  # short item names for the Prices sheet
    "mg996r": "MG996R servo", "mg90s": "MG90S micro servo", "esp32": "ESP32 DevKit", "pca9685": "PCA9685 servo driver",
    "n20": "N20 gearmotor 6 V ≈150:1", "hub1996": "Pololu #1996 hub (2-pack)", "tb6612": "TB6612FNG driver",
    "hcsr04": "HC-SR04 sensor", "brio": "Logitech Brio 4K", "psu6v5a": "6 V 5 A adapter", "dcjack": "DC jack 5.5 × 2.1",
    "dcswitch": "Inline DC switch", "fuseholder": "Inline fuse holder", "powerbus": "WAGO lever connectors + wire",
    "cells18650": "18650 cell, protected", "holder18650": "2 × 18650 holder", "charger18650": "Li-ion charger, 2 bays",
    "lipo2s": "2S LiPo ≈1000 mAh", "ubec": "UBEC 6 V ≥ 5 A", "lipocharger": "2S LiPo balance charger",
    "lipobag": "LiPo bag", "lipoalarm": "LiPo alarm", "xt30": "XT30 pigtail", "mp1584": "MP1584 buck",
    "toggle": "Mini toggle switch", "switch10a": "Switch ≥ 10 A",
    "capkit.1000uF": "1000 µF capacitor", "capkit.100uF": "100 µF capacitor", "ceramickit": "100 nF capacitors",
    "reskit.kit": "Resistor kit", "reskit.33k": "33 kΩ resistors", "schottky": "Schottky diodes", "tvs": "TVS diode",
    "breadboard": "Mini breadboard", "jumpers.mm": "Jumpers M-M", "jumpers.mf": "Jumpers M-F", "jumpers.ff": "Jumpers F-F",
    "servoext30": "Servo extension 30 cm", "servoext50": "Servo extension 60 cm", "servoext15": "Servo extension 15 cm",
    "usbc": "USB-C to USB-A cable 2 m", "ziptie": "Zip ties", "velcro": "Velcro battery strap", "heatshrink": "Heat-shrink kit",
    "wire22": "22 AWG silicone wire kit", "m3shcs": "M3 socket head kit", "m3ft50": "M3 × 50 fully threaded 12.9",
    "m3ft40": "M3 × 40 fully threaded 12.9", "m3set20": "M3 × 20 set screw", "m3nyloc": "M3 nyloc nuts",
    "m3washer.iso7089": "M3 washers Ø7", "m3washer.din9021": "M3 penny washers Ø9", "inserts": "M3 heat-set inserts",
    "dowel2x24": "Ø2 mm steel rod 100 mm", "m2kit": "M2 screws + nuts", "m2nyloc": "M2 Nyloc nut",
    "selftap": "Self-tapping screws", "m25tap": "M2.5 × 6 self-tapping", "m17": "M1.7 × 5 horn screws",
    "woodscrew": "Wood screws #4 × 2 in", "m3nut": "M3 hex nuts", "m2washer": "M2 washers", "thumbscrew": "1/4\"-20 thumb screw", "ptfe": "PTFE tape",
    "rubberfeet": "Rubber feet", "consumables.foamtape": "Mounting tape", "consumables.hotglue": "Hot-glue sticks",
    "consumables.ca": "CA glue", "consumables.threadlocker": "Threadlocker (blue)", "flatbar": "14 ga steel strip",
    "petg.generic": "PETG 1 kg (Sunlu)", "petg.prusament": "PETG 1 kg (Prusament)", "electricity": "Electricity (NB Power)",
}

# ---------------------------------------------------------------------------------------- styles
F = "Arial"
H1 = Font(name=F, size=16, bold=True)
H2 = Font(name=F, size=11, bold=True)
NORMAL = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
INPUT = Font(name=F, size=10, color="0000FF")
LINK = Font(name=F, size=10, color="008000")
MUTED = Font(name=F, size=9, color="555555", italic=True)
HEAD_FILL = PatternFill("solid", fgColor="E8ECEF")
TOTAL_FILL = PatternFill("solid", fgColor="F3F5F7")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="C9D1D8")
BOTTOM = Border(bottom=THIN)
MONEY = '"$"#,##0.00;("$"#,##0.00);"-"'
MONEY0 = '"$"#,##0;("$"#,##0);"-"'
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def head(ws, row, labels, widths=None):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = BOLD; c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="bottom"); c.border = BOTTOM
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ---------------------------------------------------------------------------------------- Prices
wp = wb.active; wp.title = "Prices"
wp["A1"] = "Prices (CAD, before tax and shipping)"; wp["A1"].font = H1
wp["A2"] = (f"Checked {AS_OF}. Blue cells are inputs you can change; everything else recalculates. USD prices convert at the "
            "rate on the Summary sheet. Confidence: high = price seen on the product page; medium = search snippet, "
            "unclear variant or a typical store price; estimate = no price could be read, a typical price is used.")
wp["A2"].font = MUTED; wp["A2"].alignment = WRAP; wp.merge_cells("A2:N2"); wp.row_dimensions[2].height = 42
cols = ["Key", "Item", "Listing", "Retailer", "Pack qty", "Pack price", "Currency", "Pack price (CAD)", "Unit price (CAD)",
        "Budget option", "Budget pack qty", "Budget pack (CAD)", "Confidence", "Notes / link"]
head(wp, 4, cols, [16, 24, 46, 24, 8, 10, 8, 12, 11, 34, 9, 12, 11, 60])
keys = sorted(main)
PR = {}
for i, k in enumerate(keys):
    r = 5 + i
    m = main[k]
    PR[k] = r
    wp.cell(row=r, column=1, value=k).font = NORMAL
    wp.cell(row=r, column=2, value=NAMES.get(k, k)).font = NORMAL
    wp.cell(row=r, column=3, value=m["listing"]).font = NORMAL
    wp.cell(row=r, column=4, value=m["retailer"]).font = NORMAL
    c = wp.cell(row=r, column=5, value=float(m["pack_qty"] or 1)); c.font = INPUT
    c = wp.cell(row=r, column=6, value=float(m["pack_price"])); c.font = INPUT; c.number_format = '#,##0.00'
    wp.cell(row=r, column=7, value=m["currency"]).font = INPUT
    c = wp.cell(row=r, column=8, value=f'=IF(G{r}="USD",F{r}*Summary!$C$8,F{r})'); c.number_format = MONEY; c.font = NORMAL
    c = wp.cell(row=r, column=9, value=f"=H{r}/E{r}"); c.number_format = MONEY; c.font = NORMAL
    b = budget.get(k)
    if b and k not in ("mp1584",):
        wp.cell(row=r, column=10, value=f'{b["listing"]} ({b["retailer"].replace(" (budget)", "")})').font = NORMAL
        c = wp.cell(row=r, column=11, value=float(b["pack_qty"] or 1)); c.font = INPUT
        bp = float(b["pack_price"])
        c = wp.cell(row=r, column=12, value=(f"={bp}*Summary!$C$8" if b["currency"] == "USD" else bp))
        c.number_format = MONEY; c.font = INPUT
    wp.cell(row=r, column=13, value=m["confidence"]).font = NORMAL
    note = m["notes"]
    if m.get("url"):
        note = (note + " · " if note else "") + m["url"]
    wp.cell(row=r, column=14, value=note).font = NORMAL
    for col in range(1, 15):
        wp.cell(row=r, column=col).alignment = WRAP if col in (3, 4, 10, 14) else TOP
LAST = 4 + len(keys)
wp.freeze_panes = "B5"
RNG = lambda col: f"Prices!${col}$5:${col}${LAST}"   # noqa: E731


def lookup(col, keycell):
    return f"INDEX({RNG(col)},MATCH({keycell},{RNG('A')},0))"


# ---------------------------------------------------------------------------------------- robot sheets
ROBOT_SHEETS = {}


def robot_sheet(name, lines, print_info):
    ws = wb.create_sheet(name)
    ws["A1"] = f"{name}: bought parts and print cost"; ws["A1"].font = H1
    kg, hours, src = print_info
    ws["A2"] = ("Quantities come from the build guide's bill of materials. 'Value of parts used' counts only the share of "
                "each pack this robot uses; 'Cost to buy' is what you pay buying whole packs for this robot alone.")
    ws["A2"].font = MUTED; ws["A2"].alignment = WRAP; ws.merge_cells("A2:L2"); ws.row_dimensions[2].height = 30
    cols = ["Group", "Item", "Price key", "Qty", "Unit price", "Value of parts used", "Pack qty", "Packs to buy",
            "Pack price", "Cost to buy", "Budget cost", "Notes"]
    head(ws, 4, cols, [16, 46, 18, 7, 10, 12, 8, 8, 10, 11, 11, 48])
    r = 5
    first = r
    for group, item, key, qty, note in lines:
        if key not in PR:
            raise SystemExit(f"{name}: no price for {key}")
        ws.cell(row=r, column=1, value=group).font = NORMAL
        ws.cell(row=r, column=2, value=item).font = NORMAL
        ws.cell(row=r, column=3, value=key).font = LINK
        c = ws.cell(row=r, column=4, value=qty); c.font = INPUT
        c = ws.cell(row=r, column=5, value=f"={lookup('I', f'C{r}')}"); c.number_format = MONEY; c.font = LINK
        c = ws.cell(row=r, column=6, value=f"=D{r}*E{r}"); c.number_format = MONEY; c.font = NORMAL
        c = ws.cell(row=r, column=7, value=f"={lookup('E', f'C{r}')}"); c.font = LINK
        c = ws.cell(row=r, column=8, value=f"=IF(D{r}<=0,0,CEILING(D{r}/G{r},1))"); c.font = NORMAL
        c = ws.cell(row=r, column=9, value=f"={lookup('H', f'C{r}')}"); c.number_format = MONEY; c.font = LINK
        c = ws.cell(row=r, column=10, value=f"=H{r}*I{r}"); c.number_format = MONEY; c.font = NORMAL
        # budget: whole budget packs when a budget option exists, else the main cost
        bq, bp = lookup('K', f'C{r}'), lookup('L', f'C{r}')
        c = ws.cell(row=r, column=11,
                    value=f"=IF(D{r}<=0,0,IF(ISNUMBER({bp}),CEILING(D{r}/{bq},1)*{bp},J{r}))")
        c.number_format = MONEY; c.font = NORMAL
        ws.cell(row=r, column=12, value=note).font = NORMAL
        for col in range(1, 13):
            ws.cell(row=r, column=col).alignment = WRAP if col in (2, 12) else TOP
        r += 1
    last = r - 1
    # printing
    r += 1
    ws.cell(row=r, column=1, value="Printing").font = H2
    r += 1
    pet = r
    ws.cell(row=pet, column=1, value="Printed").font = NORMAL
    ws.cell(row=pet, column=2, value="PETG filament (kg)").font = NORMAL
    c = ws.cell(row=pet, column=4, value=kg); c.font = INPUT
    c = ws.cell(row=pet, column=5, value="=Summary!$C$9"); c.number_format = MONEY; c.font = LINK
    c = ws.cell(row=pet, column=6, value=f"=D{pet}*E{pet}"); c.number_format = MONEY
    c = ws.cell(row=pet, column=10, value=f"=CEILING(D{pet},1)*E{pet}"); c.number_format = MONEY
    c = ws.cell(row=pet, column=11, value=f"=F{pet}"); c.number_format = MONEY
    ws.cell(row=pet, column=12, value=f"{src}. 'Cost to buy' counts whole 1 kg spools.").font = NORMAL
    ele = pet + 1
    ws.cell(row=ele, column=1, value="Printed").font = NORMAL
    ws.cell(row=ele, column=2, value="Print electricity (hours)").font = NORMAL
    c = ws.cell(row=ele, column=4, value=hours); c.font = INPUT
    c = ws.cell(row=ele, column=5, value="=Summary!$C$11*Summary!$C$10"); c.number_format = '"$"#,##0.000'; c.font = LINK
    for col in (6, 10, 11):
        c = ws.cell(row=ele, column=col, value=f"=D{ele}*E{ele}"); c.number_format = MONEY
    ws.cell(row=ele, column=12, value="Printer power × NB Power rate (Summary sheet)").font = NORMAL
    # totals
    r = ele + 2
    tot = r
    ws.cell(row=tot, column=2, value="Total before tax").font = BOLD
    for col, letter in ((6, "F"), (10, "J"), (11, "K")):
        c = ws.cell(row=tot, column=col, value=f"=SUM({letter}{first}:{letter}{last})+{letter}{pet}+{letter}{ele}")
        c.number_format = MONEY; c.font = BOLD; c.fill = TOTAL_FILL
    hst = tot + 1
    ws.cell(row=hst, column=2, value="HST").font = NORMAL
    for col, letter in ((6, "F"), (10, "J"), (11, "K")):
        c = ws.cell(row=hst, column=col, value=f"={letter}{tot}*Summary!$C$7"); c.number_format = MONEY
    grand = hst + 1
    ws.cell(row=grand, column=2, value="Total with HST").font = BOLD
    for col, letter in ((6, "F"), (10, "J"), (11, "K")):
        c = ws.cell(row=grand, column=col, value=f"={letter}{tot}+{letter}{hst}"); c.number_format = MONEY
        c.font = BOLD; c.fill = TOTAL_FILL
    ws.freeze_panes = "C5"
    ROBOT_SHEETS[name] = dict(first=first, last=last, pet=pet, ele=ele, tot=tot, grand=grand)
    return ws


# split the rover's combined supplies line
rover = [ln for ln in bom.ROVER if ln[2] != "consumables.threadlocker"] + [
    ("Supplies", "Double-sided mounting tape", "consumables.foamtape", 0.2, "Share of a roll"),
    ("Supplies", "Hot-glue sticks", "consumables.hotglue", 2, ""),
    ("Supplies", "Threadlocker, medium (blue)", "consumables.threadlocker", 0.2, "Share of a bottle"),
]
hexapod = list(bom.HEXAPOD) + [("Supplies", "Double-sided mounting tape", "consumables.foamtape", 0.2, "Share of a roll")]
robot_sheet("Arm", bom.ARM, bom.PRINT["Arm"])
robot_sheet("Rover", rover, bom.PRINT["Rover"])
robot_sheet("Hexapod", hexapod, bom.PRINT["Hexapod"])

# ---------------------------------------------------------------------------------------- all three together
wa = wb.create_sheet("All three")
wa["A1"] = "All three robots bought together"; wa["A1"].font = H1
wa["A2"] = ("Shared packs are bought once: one servo order, one resistor kit, one screw kit, and so on. Quantities add up "
            "across the three robot sheets.")
wa["A2"].font = MUTED; wa["A2"].alignment = WRAP; wa.merge_cells("A2:H2")
head(wa, 4, ["Price key", "Item", "Total qty", "Pack qty", "Packs to buy", "Pack price", "Cost to buy", "Budget cost"],
     [18, 30, 9, 8, 8, 11, 12, 12])
used = sorted({ln[2] for ln in bom.ARM + rover + hexapod})
r = 5
for k in used:
    wa.cell(row=r, column=1, value=k).font = LINK
    wa.cell(row=r, column=2, value=NAMES.get(k, k)).font = NORMAL
    sumifs = "+".join(f"SUMIF('{s}'!$C${ROBOT_SHEETS[s]['first']}:$C${ROBOT_SHEETS[s]['last']},A{r},"
                      f"'{s}'!$D${ROBOT_SHEETS[s]['first']}:$D${ROBOT_SHEETS[s]['last']})" for s in ROBOT_SHEETS)
    wa.cell(row=r, column=3, value=f"={sumifs}").font = NORMAL
    wa.cell(row=r, column=4, value=f"={lookup('E', f'A{r}')}").font = LINK
    wa.cell(row=r, column=5, value=f"=IF(C{r}<=0,0,CEILING(C{r}/D{r},1))").font = NORMAL
    c = wa.cell(row=r, column=6, value=f"={lookup('H', f'A{r}')}"); c.number_format = MONEY; c.font = LINK
    c = wa.cell(row=r, column=7, value=f"=E{r}*F{r}"); c.number_format = MONEY
    bq, bp = lookup('K', f'A{r}'), lookup('L', f'A{r}')
    c = wa.cell(row=r, column=8, value=f"=IF(C{r}<=0,0,IF(ISNUMBER({bp}),CEILING(C{r}/{bq},1)*{bp},G{r}))")
    c.number_format = MONEY
    r += 1
alast = r - 1
r += 1
apet = r
wa.cell(row=apet, column=2, value="PETG filament (kg), all three").font = NORMAL
wa.cell(row=apet, column=3, value="=" + "+".join(f"'{s}'!D{ROBOT_SHEETS[s]['pet']}" for s in ROBOT_SHEETS)).font = LINK
c = wa.cell(row=apet, column=7, value=f"=CEILING(C{apet},1)*Summary!$C$9"); c.number_format = MONEY
c = wa.cell(row=apet, column=8, value=f"=C{apet}*Summary!$C$9"); c.number_format = MONEY
aele = apet + 1
wa.cell(row=aele, column=2, value="Print electricity, all three").font = NORMAL
c = wa.cell(row=aele, column=7, value="=" + "+".join(f"'{s}'!J{ROBOT_SHEETS[s]['ele']}" for s in ROBOT_SHEETS))
c.number_format = MONEY; c.font = LINK
c = wa.cell(row=aele, column=8, value=f"=G{aele}"); c.number_format = MONEY
atot = aele + 2
wa.cell(row=atot, column=2, value="Total before tax").font = BOLD
for col, L in ((7, "G"), (8, "H")):
    c = wa.cell(row=atot, column=col, value=f"=SUM({L}5:{L}{alast})+{L}{apet}+{L}{aele}")
    c.number_format = MONEY; c.font = BOLD; c.fill = TOTAL_FILL
wa.freeze_panes = "B5"

# ---------------------------------------------------------------------------------------- Summary (first sheet)
ws = wb.create_sheet("Summary", 0)
ws["A1"] = "Robot build costs"; ws["A1"].font = H1
ws["A2"] = (f"Desktop robot arm, tracked rover and hexapod. Canadian dollars, prices checked {AS_OF} at Canadian "
            "retailers (some US stores where nothing Canadian was found). Shipping is not included. Change any blue "
            "cell and every total updates.")
ws["A2"].font = MUTED; ws["A2"].alignment = WRAP; ws.merge_cells("A2:H2"); ws.row_dimensions[2].height = 30
ws["A5"] = "Assumptions"; ws["A5"].font = H2
assum = [
    (7, "HST (New Brunswick)", 0.15, "0%", "Harmonized sales tax"),
    (8, "USD to CAD", FX, "0.0000", "Bank of Canada, 23 Oct 2026"),
    (9, "PETG price per kg", "=Prices!H%d" % PR["petg.generic"], MONEY, "Sunlu PETG 1 kg at Canada Computers; Prusament PETG is on the Prices sheet"),
    (10, "Electricity ($/kWh)", "=Prices!H%d" % PR["electricity"], '"$"0.0000', "NB Power residential energy charge"),
    (11, "Printer power while printing PETG (kW)", 0.09, "0.00", "Assumption: MK3S average with the bed at 80-90 °C"),
]
for row, label, val, fmt, note in assum:
    ws.cell(row=row, column=2, value=label).font = NORMAL
    c = ws.cell(row=row, column=3, value=val); c.number_format = fmt
    c.font = LINK if isinstance(val, str) else INPUT
    ws.cell(row=row, column=4, value=note).font = MUTED
ws["A13"] = "What each robot costs"; ws["A13"].font = H2
head(ws, 14, ["", "Robot", "Bought parts", "PETG", "Electricity", "Total before tax", "Total with HST",
              "Budget total with HST"], [4, 34, 13, 10, 11, 15, 15, 18])
ws.column_dimensions["D"].width = 12
rows = [("Desktop robot arm", "Arm"), ("Tracked rover", "Rover"), ("Hexapod", "Hexapod")]
r = 15
for label, s in rows:
    R = ROBOT_SHEETS[s]
    ws.cell(row=r, column=2, value=label).font = NORMAL
    c = ws.cell(row=r, column=3, value=f"=SUM('{s}'!J{R['first']}:J{R['last']})"); c.number_format = MONEY0; c.font = LINK
    c = ws.cell(row=r, column=4, value=f"='{s}'!J{R['pet']}"); c.number_format = MONEY0; c.font = LINK
    c = ws.cell(row=r, column=5, value=f"='{s}'!J{R['ele']}"); c.number_format = MONEY; c.font = LINK
    c = ws.cell(row=r, column=6, value=f"=SUM(C{r}:E{r})"); c.number_format = MONEY0
    c = ws.cell(row=r, column=7, value=f"=F{r}*(1+$C$7)"); c.number_format = MONEY0; c.font = BOLD
    c = ws.cell(row=r, column=8, value=f"='{s}'!K{R['tot']}*(1+$C$7)"); c.number_format = MONEY0
    r += 1
ws.cell(row=r, column=2, value="All three, bought together").font = BOLD
c = ws.cell(row=r, column=3, value=f"=SUM('All three'!G5:G{alast})"); c.number_format = MONEY0; c.font = LINK
c = ws.cell(row=r, column=4, value=f"='All three'!G{apet}"); c.number_format = MONEY0; c.font = LINK
c = ws.cell(row=r, column=5, value=f"='All three'!G{aele}"); c.number_format = MONEY; c.font = LINK
c = ws.cell(row=r, column=6, value=f"=SUM(C{r}:E{r})"); c.number_format = MONEY0; c.font = BOLD
c = ws.cell(row=r, column=7, value=f"=F{r}*(1+$C$7)"); c.number_format = MONEY0; c.font = BOLD
c = ws.cell(row=r, column=8, value=f"='All three'!H{atot}*(1+$C$7)"); c.number_format = MONEY0; c.font = BOLD
for col in range(2, 9):
    ws.cell(row=r, column=col).fill = TOTAL_FILL
together = r
r += 1
ws.cell(row=r, column=2, value="Saving from shared packs").font = NORMAL
c = ws.cell(row=r, column=7, value=f"=SUM(G15:G17)-G{together}"); c.number_format = MONEY0
ws.cell(row=r + 2, column=2, value="Value of the parts each robot actually uses (its share of every pack)").font = H2
r += 3
head(ws, r, ["", "Robot", "Parts used", "PETG used", "Electricity", "Total before tax", "Total with HST", ""])
r += 1
for label, s in rows:
    R = ROBOT_SHEETS[s]
    ws.cell(row=r, column=2, value=label).font = NORMAL
    c = ws.cell(row=r, column=3, value=f"=SUM('{s}'!F{R['first']}:F{R['last']})"); c.number_format = MONEY0; c.font = LINK
    c = ws.cell(row=r, column=4, value=f"='{s}'!F{R['pet']}"); c.number_format = MONEY0; c.font = LINK
    c = ws.cell(row=r, column=5, value=f"='{s}'!F{R['ele']}"); c.number_format = MONEY; c.font = LINK
    c = ws.cell(row=r, column=6, value=f"=SUM(C{r}:E{r})"); c.number_format = MONEY0
    c = ws.cell(row=r, column=7, value=f"=F{r}*(1+$C$7)"); c.number_format = MONEY0
    r += 1
r += 1
notes = [
    "Not counted: the Logitech Brio 4K you already own (C$249.99 new at Canada Computers), shipping, and tools.",
    "Tools you may need: soldering iron (also sets the rover's heat-set inserts), hex keys 1.5 and 2.5 mm, small Phillips "
    "screwdriver, multimeter, calipers, needle file, hacksaw and file for the rover's steel plates.",
    "Budget total: the same build with the cheapest price found per item (servos direct from Waveshare, and so on).",
    "Prices with confidence 'estimate' on the Prices sheet are typical store prices where no listing could be read; "
    "check them before ordering. Amazon.ca was not readable by the price search and is often cheaper for multi-packs.",
    "PETG and print hours: arm from its latest slices; rover from its round-2 slices; hexapod estimated. They change "
    "with each design round.",
]
for n in notes:
    ws.cell(row=r, column=2, value="• " + n).font = NORMAL
    ws.cell(row=r, column=2).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 42
    r += 1
ws.freeze_panes = "A5"

for sh in wb.worksheets:
    sh.page_setup.orientation = "landscape"
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.print_options.gridLines = False
wb.save(OUT)
print("saved", OUT)
